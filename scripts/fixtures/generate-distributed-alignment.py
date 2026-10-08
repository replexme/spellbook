# SPDX-License-Identifier: MPL-2.0
# Local regression fixture; requires python-pptx. No customer material.
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR

cases = [
    ('digits', '319', 'Arial', 'dist', False),
    ('korean', '한글 정렬', 'Malgun Gothic', 'dist', False),
    ('spaces', 'A B C', 'Arial', 'dist', False),
    ('single', 'A', 'Arial', 'dist', False),
    ('hebrew', 'שלום עולם', 'Arial', 'dist', True),
    ('cjk', '中文对齐', 'Arial', 'dist', False),
    ('last-line', 'The final line\nspreads too', 'Arial', 'dist', False),
    ('wrapped', 'Long words with spaces wrap over multiple lines and distribute even the final line', 'Arial', 'dist', False),
    ('justify-control', 'A B C', 'Arial', 'just', False),
    ('hebrew-single', 'שלום', 'Arial', 'dist', True),
    ('arabic', 'مرحبا بالعالم', 'Arial', 'dist', True),
    ('arabic-single', 'مرحبا', 'Arial', 'dist', True),
]

presentation = Presentation()
presentation.slide_width = Inches(13.333333)
presentation.slide_height = Inches(7.5)
for name, text, family, alignment, rtl in cases:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Inches(.1)
    frame.margin_right = Inches(.1)
    for index, line in enumerate(text.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        properties = paragraph._p.get_or_add_pPr()
        properties.set('algn', alignment)
        if rtl:
            properties.set('rtl', '1')
        for run in paragraph.runs:
            run.font.name = family
            run.font.size = Pt(32 if name == 'wrapped' else 44)
    slide.notes_slide.notes_text_frame.text = 'Distributed-alignment regression: ' + name
presentation.save(Path(__file__).resolve().parents[2] / 'eval/public/fixtures/distributed-paragraph-alignment.pptx')
